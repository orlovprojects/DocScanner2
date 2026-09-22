"""xlsx failų skaitymas į OpeningBalanceLine objektus (dar neišsaugotus)."""

import unicodedata
from decimal import Decimal, InvalidOperation
import re

from openpyxl import load_workbook

from ..models import OpeningBalanceLine, OpeningBalanceSection

MONEY = Decimal("0.01")


class ParseError(Exception):
    pass


def _strip_diacritics(text):
    normalized = unicodedata.normalize("NFKD", str(text or ""))
    return "".join(c for c in normalized if not unicodedata.combining(c))


def _to_decimal(value):
    if value in (None, ""):
        return Decimal("0")
    if isinstance(value, (int, float, Decimal)):
        raw = str(value)
    else:
        raw = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return Decimal(raw).quantize(MONEY)
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _text(value, limit=255):
    return str(value or "").strip()[:limit]


# Stulpelių atpažinimas. Aliasai be diakritikos ir mažosiomis.
# Tvarka svarbi: konkretesni variantai tikrinami pirmiau.
# Ilgalaikio turto failas
FIXED_ASSET_ALIASES = {
    "inventory_number": ("inventorinis nr", "inventorinis numeris", "inv nr", "inventorinis"),
    "name": ("pavadinimas", "turto pavadinimas"),
    "category": ("turto grupe", "grupe", "kategorija"),
    "purchase_date": ("isigijimo data", "pirkimo data"),
    "operation_start_date": ("eksploatacijos pradzia", "naudojimo pradzia"),
    "acquisition_cost": ("isigijimo savikaina", "savikaina", "verte"),
    "accumulated": ("sukauptas nusidevejimas", "nusidevejimas", "amortizacija"),
    "useful_life_months": ("naudingo tarnavimo laikas", "tarnavimo laikas", "menesiai", "laikotarpis"),
    "salvage_value": ("likvidacine verte", "likvidacine"),
}


HEADER_ALIASES = {
    # Kontrahentų failai
    "debt": ("pirkejo skola", "skola tiekejui", "skola"),
    "advance": ("pirkejo permoka", "avansas tiekejui", "permoka", "avansas"),
    "counterparty_code": ("imones kodas", "kontrahento kodas", "kliento kodas", "kodas"),
    "counterparty_vat_code": ("pvm kodas", "pvm", "vat"),
    "counterparty_name": (
        "pavadinimas / vardas pavarde", "pavadinimas/vardas pavarde",
        "kontrahentas", "pirkejas", "tiekejas", "klientas", "imone",
    ),
    "address": ("adresas", "address"),
    "country": ("salis", "country"),
    "person_type": ("tipas", "type"),
    # Balansas / bankai
    "amount_currency": ("suma valiuta", "amount currency", "valiutos suma"),
    "account_name": ("saskaitos pavadinimas", "pavadinimas", "name"),
    "account_code": ("saskaita", "account", "code"),
    "bank_name": ("banko pavadinimas", "bankas", "bank"),
    "iban": ("iban", "saskaitos nr", "banko saskaita"),
    "debit": ("debetas", "debit"),
    "credit": ("kreditas", "credit"),
    "currency": ("valiuta", "currency"),
}

# Kontrahentų sekcijų sąskaitos
COUNTERPARTY_ACCOUNTS = {
    OpeningBalanceSection.BUYER: {"debt": "2410", "advance": "4420"},
    OpeningBalanceSection.SUPPLIER: {"debt": "4430", "advance": "2080"},
}


def _detect_fixed_asset_columns(header_row):
    found = {}
    for idx, cell in enumerate(header_row):
        title = _strip_diacritics(cell).strip().lower()
        if not title:
            continue
        for field, aliases in FIXED_ASSET_ALIASES.items():
            if field in found:
                continue
            if any(title == a or title.startswith(a) for a in aliases):
                found[field] = idx
                break
    return found


def _to_date(value):
    import datetime

    if value in (None, ""):
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value

    raw = str(value).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%Y.%m.%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None


def parse_fixed_assets(file_obj, categories):
    """
    Grąžina (lines, warnings). categories — {normalizuotas pavadinimas: kodas}.
    Duomenys saugomi OpeningBalanceLine.extra.
    """
    rows = _read_rows(file_obj)

    header_idx, cols = None, None
    for idx, row in enumerate(rows[:10]):
        found = _detect_fixed_asset_columns(row)
        if "name" in found and "acquisition_cost" in found:
            header_idx, cols = idx, found
            break

    if cols is None:
        raise ParseError("Nerasti stulpeliai „Pavadinimas“ ir „Įsigijimo savikaina“.")

    def cell(row, field):
        i = cols.get(field)
        return row[i] if i is not None and i < len(row) else None

    lines = []
    warnings = []
    seen_numbers = set()

    for offset, row in enumerate(rows[header_idx + 1:], start=1):
        if not any(str(c or "").strip() for c in row):
            continue

        name = _text(cell(row, "name"))
        inventory_number = _text(cell(row, "inventory_number"), 64)

        if _is_total_row(name, inventory_number):
            continue

        cost = abs(_to_decimal(cell(row, "acquisition_cost")))
        if not name and cost == 0:
            continue

        if not name:
            warnings.append(f"{offset} eilutė: nenurodytas pavadinimas, praleista.")
            continue

        if cost <= 0:
            warnings.append(f"{offset} eilutė ({name}): nenurodyta savikaina, praleista.")
            continue

        if inventory_number:
            if inventory_number in seen_numbers:
                warnings.append(f"{offset} eilutė ({name}): kartojasi inventorinis nr. {inventory_number}.")
            seen_numbers.add(inventory_number)

        raw_category = _strip_diacritics(cell(row, "category")).strip().lower()
        category = categories.get(raw_category, "")
        if raw_category and not category:
            warnings.append(f"{offset} eilutė ({name}): nežinoma grupė „{_text(cell(row, 'category'))}“.")

        accumulated = abs(_to_decimal(cell(row, "accumulated")))
        salvage = abs(_to_decimal(cell(row, "salvage_value")))

        if accumulated > cost - salvage:
            warnings.append(
                f"{offset} eilutė ({name}): sukauptas nusidėvėjimas viršija nudėvimą vertę."
            )

        life = cell(row, "useful_life_months")
        try:
            life = int(_to_decimal(life)) if life not in (None, "") else 0
        except (ValueError, TypeError):
            life = 0

        purchase_date = _to_date(cell(row, "purchase_date"))
        start_date = _to_date(cell(row, "operation_start_date")) or purchase_date

        if not purchase_date:
            warnings.append(f"{offset} eilutė ({name}): nenurodyta įsigijimo data.")

        line = OpeningBalanceLine(
            section=OpeningBalanceSection.FIXED_ASSET,
            account_code=category,
            account_name=name,
            debit=cost,
            credit=accumulated,
            row_number=offset,
            extra={
                "inventory_number": inventory_number,
                "category": category,
                "raw_category": _text(cell(row, "category")),
                "purchase_date": purchase_date.isoformat() if purchase_date else "",
                "operation_start_date": start_date.isoformat() if start_date else "",
                "useful_life_months": life,
                "salvage_value": str(salvage),
            },
        )
        lines.append(line)

    if not lines:
        raise ParseError("Faile nerasta turto eilučių.")

    return lines, warnings


def _detect_columns(header_row, section):
    """Grąžina {laukas: stulpelio_indeksas}."""
    is_counterparty = section in COUNTERPARTY_ACCOUNTS

    found = {}
    for idx, cell in enumerate(header_row):
        title = _strip_diacritics(cell).strip().lower()
        if not title:
            continue
        for field, aliases in HEADER_ALIASES.items():
            if field in found:
                continue
            # Kontrahentų failuose „pavadinimas" yra kontrahentas, ne sąskaita
            if is_counterparty and field in ("account_name", "account_code"):
                continue
            if not is_counterparty and field in ("debt", "advance", "counterparty_name"):
                continue
            if any(title == a or title.startswith(a + " ") for a in aliases):
                found[field] = idx
                break
    return found


def _read_rows(file_obj):
    wb = load_workbook(file_obj, data_only=True, read_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not rows:
        raise ParseError("Failas tuščias.")
    return rows


def _find_header(rows, section):
    """Randa antraštės eilutę per pirmas 10 eilučių."""
    required = (
        ("counterparty_name",)
        if section in COUNTERPARTY_ACCOUNTS
        else ("debit", "credit")
    )
    for idx, row in enumerate(rows[:10]):
        cols = _detect_columns(row, section)
        if all(r in cols for r in required):
            return idx, cols

    if section in COUNTERPARTY_ACCOUNTS:
        raise ParseError("Nerastas kontrahento pavadinimo stulpelis.")
    raise ParseError("Nerasti stulpeliai „Debetas“ ir „Kreditas“.")


def _is_total_row(name, code):
    if code:
        return False
    lowered = _strip_diacritics(name).strip().lower()
    return lowered.startswith(("is viso", "viso", "total", "suma"))


def parse_section(file_obj, section):
    """
    Grąžina (lines, warnings).
    lines — neišsaugoti OpeningBalanceLine be batch.
    """
    rows = _read_rows(file_obj)
    header_idx, cols = _find_header(rows, section)

    def cell(row, field):
        idx = cols.get(field)
        return row[idx] if idx is not None and idx < len(row) else None

    if section in COUNTERPARTY_ACCOUNTS:
        return _parse_counterparty(rows, header_idx, cell, section)
    return _parse_accounts(rows, header_idx, cell, section)


def _parse_counterparty(rows, header_idx, cell, section):
    """Pirkėjų arba tiekėjų skolų failas."""
    accounts = COUNTERPARTY_ACCOUNTS[section]
    lines = []
    warnings = []
    seen = {}

    for offset, row in enumerate(rows[header_idx + 1:], start=1):
        if not any(str(c or "").strip() for c in row):
            continue

        name = _text(cell(row, "counterparty_name"))
        code = _text(cell(row, "counterparty_code"), 50)

        if _is_total_row(name, code):
            continue

        debt = abs(_to_decimal(cell(row, "debt")))
        advance = abs(_to_decimal(cell(row, "advance")))

        if debt == 0 and advance == 0:
            continue

        if not name and not code:
            warnings.append(f"{offset} eilutė: nenurodytas kontrahentas, praleista.")
            continue

        currency = (_text(cell(row, "currency"), 3) or "EUR").upper()
        vat_code = _text(cell(row, "counterparty_vat_code"), 50)
        address = _text(cell(row, "address"))
        country = _text(cell(row, "country"), 50)
        person_type = _text(cell(row, "person_type"), 50)
        ibans = [
            part.strip().replace(" ", "").upper()
            for part in re.split(r"[;,\n]", _text(cell(row, "iban"), 512))
            if part.strip()
        ]

        if debt and advance:
            warnings.append(
                f"{offset} eilutė ({name}): užpildyta ir skola, ir permoka — "
                f"įrašomos abi sumos."
            )

        for kind, amount in (("debt", debt), ("advance", advance)):
            if amount == 0:
                continue

            account = accounts[kind]
            key = (code or name.upper(), kind, currency)
            if key in seen:
                warnings.append(
                    f"{offset} eilutė ({name}): kartojasi su ta pačia valiuta — "
                    f"sumos sudedamos."
                )
            seen[key] = True

            line = OpeningBalanceLine(
                section=section,
                account_code=account,
                mapped_account=account,
                match_type=OpeningBalanceLine.MATCH_EXACT,
                counterparty_name=name,
                counterparty_code=code,
                counterparty_vat_code=vat_code,
                currency=currency,
                row_number=offset,
                extra={
                    "kind": kind,
                    "address": address,
                    "country": country,
                    "person_type": person_type,
                    "ibans": ibans,
                },
            )

            # Suma visada teigiama; pusę nulemia sąskaita
            if account in ("2410", "2080"):
                line.debit = amount
            else:
                line.credit = amount

            if currency != "EUR":
                line.amount_currency = amount

            lines.append(line)

    if not lines:
        raise ParseError("Faile nerasta eilučių su sumomis.")

    return lines, warnings


def _parse_accounts(rows, header_idx, cell, section):
    """Balanso arba banko sąskaitų failas."""
    lines = []
    warnings = []

    for offset, row in enumerate(rows[header_idx + 1:], start=1):
        if not any(str(c or "").strip() for c in row):
            continue

        code = _text(cell(row, "account_code"), 40)
        name = _text(cell(row, "account_name"))

        if _is_total_row(name, code):
            continue

        debit = _to_decimal(cell(row, "debit"))
        credit = _to_decimal(cell(row, "credit"))

        if debit == 0 and credit == 0:
            continue

        if debit != 0 and credit != 0:
            warnings.append(f"{offset} eilutė: užpildyti ir debetas, ir kreditas.")

        if debit < 0 or credit < 0:
            warnings.append(f"{offset} eilutė: neigiama suma perkelta į kitą pusę.")
            if debit < 0:
                credit += abs(debit)
                debit = Decimal("0")
            if credit < 0:
                debit += abs(credit)
                credit = Decimal("0")

        line = OpeningBalanceLine(
            section=section,
            account_code=code,
            account_name=name,
            debit=debit,
            credit=credit,
            currency=(_text(cell(row, "currency"), 3) or "EUR").upper(),
            row_number=offset,
        )

        amount_currency = cell(row, "amount_currency")
        if amount_currency not in (None, ""):
            line.amount_currency = _to_decimal(amount_currency)

        if section == OpeningBalanceSection.BANK:
            line.extra = {
                "iban": _text(cell(row, "iban"), 64).replace(" ", "").upper(),
                "bank_name": _text(cell(row, "bank_name")),
            }

        lines.append(line)

    if not lines:
        raise ParseError("Faile nerasta eilučių su sumomis.")

    return lines, warnings