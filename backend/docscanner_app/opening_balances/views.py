"""Pradinių likučių API."""

from datetime import date

from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from ..models import (
    CompanyProfile,
    OpeningBalanceBatch,
    OpeningBalanceLine,
    OpeningBalanceSection,
)
from ..utils.chart_of_accounts import get_account_name, is_valid_account
from . import services
from .matching import apply_matching, save_mapping
from .parsers import ParseError, parse_section

SECTION_CODES = {s[0] for s in OpeningBalanceSection.SECTION_CHOICES}


def _profile(request):
    company_id = request.query_params.get("company_profile") or request.data.get("company_profile")
    if company_id:
        return CompanyProfile.objects.filter(id=company_id, user=request.user).first()
    return getattr(request.user, "active_company_profile", None)


def _line_payload(line):
    return {
        "id": line.id,
        "section": line.section,
        "row_number": line.row_number,
        "account_code": line.account_code,
        "account_name": line.account_name,
        "debit": str(line.debit),
        "credit": str(line.credit),
        "currency": line.currency,
        "amount_currency": str(line.amount_currency) if line.amount_currency is not None else None,
        "mapped_account": line.mapped_account,
        "mapped_name": line.mapped_name,
        "match_type": line.match_type,
        "counterparty_name": line.counterparty_name,
        "counterparty_code": line.counterparty_code,
        "counterparty_vat_code": line.counterparty_vat_code,
        "extra": line.extra or {},
        "error": line.error,
    }


def _batch_payload(batch):
    sections = {
        s.section: {
            "file_name": s.file_name,
            "uploaded_at": s.uploaded_at.isoformat() if s.uploaded_at else None,
            "line_count": batch.lines.filter(section=s.section).count(),
        }
        for s in batch.sections.all()
    }
    return {
        "id": batch.id,
        "cutover_date": batch.cutover_date.isoformat(),
        "entry_date": batch.entry_date.isoformat(),
        "status": batch.status,
        "diff_policy": batch.diff_policy,
        "journal_entry_id": batch.journal_entry_id,
        "confirmed_at": batch.confirmed_at.isoformat() if batch.confirmed_at else None,
        "sections": sections,
    }


@api_view(["GET", "POST", "DELETE"])
@permission_classes([IsAuthenticated])
def opening_batch(request):
    """
    GET    — partijos būsena (arba null).
    POST   — sukuria arba atnaujina perėjimo datą. Body: {"cutover_date": "2026-01-01"}
    DELETE — ištrina visą partiją (tik juodraštį).
    """
    profile = _profile(request)
    if not profile:
        return Response({"detail": "Nepasirinktas įmonės profilis."}, status=400)

    batch = services.get_batch(profile)

    if request.method == "GET":
        if not batch:
            return Response({"batch": None})
        return Response({"batch": _batch_payload(batch), "summary": services.reconcile(batch)})

    if request.method == "DELETE":
        if not batch:
            return Response(status=204)
        if batch.status != OpeningBalanceBatch.STATUS_DRAFT:
            return Response({"detail": "Patvirtintos partijos trinti negalima. Pirmiausia ją atšaukite."}, status=400)
        batch.delete()
        return Response(status=204)

    raw_date = str(request.data.get("cutover_date") or "").strip()
    try:
        cutover = date.fromisoformat(raw_date)
    except ValueError:
        return Response({"detail": "Nurodykite tinkamą perėjimo datą."}, status=400)

    diff_policy = request.data.get("diff_policy")

    if batch:
        if batch.status != OpeningBalanceBatch.STATUS_DRAFT:
            return Response({"detail": "Likučiai jau patvirtinti. Pirmiausia juos atšaukite."}, status=400)
        batch.cutover_date = cutover
        if diff_policy in dict(OpeningBalanceBatch.DIFF_CHOICES):
            batch.diff_policy = diff_policy
        batch.save(update_fields=["cutover_date", "diff_policy", "updated_at"])
    else:
        batch = OpeningBalanceBatch.objects.create(
            company_profile=profile,
            cutover_date=cutover,
            created_by=request.user,
            diff_policy=diff_policy if diff_policy in dict(OpeningBalanceBatch.DIFF_CHOICES) else OpeningBalanceBatch.DIFF_ASK,
        )

    return Response({"batch": _batch_payload(batch), "summary": services.reconcile(batch)})


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def opening_upload(request, section):
    """
    POST /api/apskaita/pradiniai-likuciai/upload/<section>/
    multipart: file=<xlsx>. Pakartotinis įkėlimas pakeičia visą sekciją.
    """
    if section not in SECTION_CODES:
        return Response({"detail": "Nežinoma sekcija."}, status=400)

    profile = _profile(request)
    if not profile:
        return Response({"detail": "Nepasirinktas įmonės profilis."}, status=400)

    batch = services.get_batch(profile)
    if not batch:
        return Response({"detail": "Pirmiausia nurodykite perėjimo datą."}, status=400)
    if not batch.is_editable:
        return Response({"detail": "Likučiai patvirtinti. Pirmiausia juos atšaukite."}, status=400)

    upload = request.FILES.get("file")
    if not upload:
        return Response({"detail": "Nepateiktas failas."}, status=400)
    if not upload.name.lower().endswith((".xlsx", ".xlsm")):
        return Response({"detail": "Įkelkite .xlsx failą."}, status=400)

    try:
        lines, warnings = parse_section(upload, section)
    except ParseError as e:
        return Response({"detail": str(e)}, status=400)
    except Exception as e:
        return Response({"detail": f"Nepavyko perskaityti failo: {e}"}, status=400)

    if section in (OpeningBalanceSection.BALANCE, OpeningBalanceSection.BANK):
        apply_matching(lines, profile.id)

    batch.lines.filter(section=section).delete()
    for line in lines:
        line.batch = batch
    OpeningBalanceLine.objects.bulk_create(lines)

    from django.utils import timezone
    OpeningBalanceSection.objects.update_or_create(
        batch=batch,
        section=section,
        defaults={"file_name": upload.name[:255], "uploaded_at": timezone.now()},
    )

    saved = batch.lines.filter(section=section)
    return Response({
        "section": section,
        "warnings": warnings,
        "lines": [_line_payload(l) for l in saved],
        "summary": services.reconcile(batch),
    })


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def opening_lines(request, section):
    """GET — vienos sekcijos eilutės."""
    if section not in SECTION_CODES:
        return Response({"detail": "Nežinoma sekcija."}, status=400)

    profile = _profile(request)
    batch = services.get_batch(profile) if profile else None
    if not batch:
        return Response({"lines": []})

    return Response({
        "section": section,
        "lines": [_line_payload(l) for l in batch.lines.filter(section=section)],
    })


@api_view(["PATCH"])
@permission_classes([IsAuthenticated])
def opening_line_detail(request, pk):
    """
    PATCH — rankinis sąskaitos priskyrimas eilutei.
    Body: {"mapped_account": "2410"}
    """
    profile = _profile(request)
    if not profile:
        return Response({"detail": "Nepasirinktas įmonės profilis."}, status=400)

    line = OpeningBalanceLine.objects.filter(pk=pk, batch__company_profile=profile).select_related("batch").first()
    if not line:
        return Response({"detail": "Eilutė nerasta."}, status=404)
    if not line.batch.is_editable:
        return Response({"detail": "Likučiai patvirtinti."}, status=400)

    code = str(request.data.get("mapped_account") or "").strip()
    if not code:
        return Response({"detail": "Nurodykite sąskaitą."}, status=400)
    from .matching import is_postable

    if not is_valid_account(code):
        return Response({"detail": f"Nežinoma sąskaita: {code}."}, status=400)
    if not is_postable(code):
        return Response(
            {"detail": f"{code} yra grupinė sąskaita — pasirinkite konkrečią."},
            status=400,
        )

    line.mapped_account = code
    line.mapped_name = get_account_name(code) or ""
    line.match_type = OpeningBalanceLine.MATCH_MANUAL
    line.save(update_fields=["mapped_account", "mapped_name", "match_type"])

    if line.account_code:
        save_mapping(profile.id, line.account_code, line.account_name, code)

    return Response({"line": _line_payload(line), "summary": services.reconcile(line.batch)})


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def opening_summary(request):
    """GET — patikrinimo suvestinė."""
    profile = _profile(request)
    batch = services.get_batch(profile) if profile else None
    if not batch:
        return Response({"summary": None})
    return Response({"summary": services.reconcile(batch)})


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def opening_confirm(request):
    """POST — patvirtina likučius ir sukuria DK įrašą."""
    profile = _profile(request)
    if not profile:
        return Response({"detail": "Nepasirinktas įmonės profilis."}, status=400)

    batch = services.get_batch(profile)
    if not batch:
        return Response({"detail": "Likučių partija nerasta."}, status=404)
    if batch.status == OpeningBalanceBatch.STATUS_CONFIRMED:
        return Response({"detail": "Likučiai jau patvirtinti."}, status=400)

    policy = request.data.get("diff_policy")
    if policy in dict(OpeningBalanceBatch.DIFF_CHOICES):
        batch.diff_policy = policy
        batch.save(update_fields=["diff_policy", "updated_at"])

    entry, summary = services.confirm(batch, request.user)
    if not entry:
        return Response({"detail": "Likučiai dar neparuošti patvirtinimui.", "summary": summary}, status=400)

    batch.refresh_from_db()
    return Response({"batch": _batch_payload(batch), "summary": summary, "journal_entry_id": entry.id})


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def opening_reopen(request):
    """POST — grąžina į juodraštį ir ištrina DK įrašą."""
    profile = _profile(request)
    if not profile:
        return Response({"detail": "Nepasirinktas įmonės profilis."}, status=400)

    batch = services.get_batch(profile)
    if not batch:
        return Response({"detail": "Likučių partija nerasta."}, status=404)

    ok, error = services.reopen(batch)
    if not ok:
        return Response({"detail": error}, status=400)

    batch.refresh_from_db()
    return Response({"batch": _batch_payload(batch), "summary": services.reconcile(batch)})


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def opening_template(request, section):
    """GET — tuščias xlsx šablonas."""
    import io

    from django.http import HttpResponse
    from openpyxl import Workbook

    if section not in SECTION_CODES:
        return Response({"detail": "Nežinoma sekcija."}, status=400)

    headers = {
        OpeningBalanceSection.BALANCE: [
            ["Sąskaita", "Pavadinimas", "Debetas", "Kreditas", "Valiuta", "Suma valiuta"],
            ["1230", "Transporto priemonių įsigijimo savikaina", 15000, None, None, None],
            ["2410", "Pirkėjų skolų vertė", 1250.50, None, None, None],
            ["2711", "Swedbank", 2079.50, None, None, None],
            ["3011", "Paprastosios akcijos", None, 2500, None, None],
            ["3411", "Ataskaitinių metų grynasis pelnas", None, 15000, None, None],
            ["4430", "Skolos tiekėjams už prekes ir paslaugas", None, 830, None, None],
        ],
        OpeningBalanceSection.BANK: [
            ["Sąskaita", "Pavadinimas", "IBAN", "Bankas", "Debetas", "Kreditas", "Valiuta", "Suma valiuta"],
            ["2711", "Swedbank EUR", "LT121000011101001000", "Swedbank", 12300, None, "EUR", None],
            ["2712", "Revolut USD", "LT303250000000000001", "Revolut", None, None, "USD", 1000],
        ],
        OpeningBalanceSection.BUYER: [
            [
                "Pavadinimas / Vardas Pavardė", "Kodas", "PVM kodas",
                "Pirkėjo skola", "Pirkėjo permoka", "Valiuta",
                "Adresas (nebūtina)", "Šalies kodas (nebūtina)", "Tipas (fizinis/juridinis) (nebūtina)", "IBAN (nebūtina)",
            ],
            ["UAB Pavyzdys", "300123456", "LT100001234567", 1250.50, None, "EUR", None, None, None,
             "LT121000011101001000, LT303250000000000001"],
        ],
        OpeningBalanceSection.SUPPLIER: [
            [
                "Pavadinimas / Vardas Pavardė", "Kodas", "PVM kodas",
                "Skola tiekėjui", "Avansas tiekėjui", "Valiuta",
                "Adresas (nebūtina)", "Šalies kodas (nebūtina)", "Tipas (fizinis/juridinis) (nebūtina)", "IBAN (nebūtina)",
            ],
            ["UAB Tiekėjas", "300654321", None, 830, None, "EUR", None, None, None, None],
            ["ACME Inc", None, None, None, 1000, "USD", None, "US", None, None],
        ],
    }[section]

    from openpyxl.styles import Alignment, Font, PatternFill

    widths = {
        OpeningBalanceSection.BALANCE: [12, 42, 14, 14, 10, 14],
        OpeningBalanceSection.BANK: [12, 24, 26, 16, 14, 14, 10, 14],
        OpeningBalanceSection.BUYER: [34, 14, 18, 15, 17, 10, 22, 20, 30, 42],
        OpeningBalanceSection.SUPPLIER: [34, 14, 18, 16, 18, 10, 22, 20, 30, 42],
    }[section]

    sheet_titles = {
        OpeningBalanceSection.BALANCE: "Balansas",
        OpeningBalanceSection.BANK: "Banko_saskaitos",
        OpeningBalanceSection.BUYER: "Pirkeju_skolos_permokos",
        OpeningBalanceSection.SUPPLIER: "Skolos_avansai_tiekejams",
    }

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_titles[section]
    for row in headers:
        ws.append(row)

    required_font = Font(bold=True, color="1F2937")
    optional_font = Font(bold=False, color="6B7280", italic=True)
    required_fill = PatternFill("solid", fgColor="EFF6FF")
    optional_fill = PatternFill("solid", fgColor="F9FAFB")

    for idx, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=idx)
        ws.column_dimensions[cell.column_letter].width = width

        is_optional = "nebūtina" in str(cell.value or "").lower()
        cell.font = optional_font if is_optional else required_font
        cell.fill = optional_fill if is_optional else required_fill
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    file_names = {
        OpeningBalanceSection.BALANCE: "balansas",
        OpeningBalanceSection.BANK: "banko_saskaitos",
        OpeningBalanceSection.BUYER: "pirkejai",
        OpeningBalanceSection.SUPPLIER: "tiekejai",
    }
    response["Content-Disposition"] = (
        f'attachment; filename="pradiniai_likuciai_{file_names[section]}.xlsx"'
    )
    return response